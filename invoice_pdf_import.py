#!/usr/bin/env python3
"""
EWI PDF INVOICE IMPORTER - TEST VERSION

What it does:
1. Reads a text-based PDF invoice.
2. Extracts supplier, invoice number, date, material, quantity and unit price.
3. Matches material by Material ID first, then exact material name.
4. Updates MATERIAL_PRICES:
       Previous Price = old New Price
       New Price      = invoice Unit Price
5. Adds invoice data to INVOICES / INVOICE_ITEMS / PRICE_HISTORY.
6. NEVER silently adds an unknown material - it reports it for manual review.

Mac setup:
    python3 -m pip install openpyxl pypdf

Run:
    python3 invoice_pdf_import.py Test_Invoice_Supplier_A_1001.pdf WarehouseFlow_EWI_With_Prices_And_Invoices.xlsx

IMPORTANT:
This parser is designed for the TEST PDFs we created. Real supplier PDFs can have
different layouts, so each supplier may need a small parser adjustment later.
"""

import re
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pypdf import PdfReader


def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def norm(s):
    return re.sub(r"[^a-z0-9]+", "", clean(s).lower())


def money(s):
    s = clean(s).replace("€", "").replace(",", "")
    return float(s)


def read_pdf(pdf_file):
    text = "\n".join((p.extract_text() or "") for p in PdfReader(pdf_file).pages)
    lines = [clean(x) for x in text.splitlines() if clean(x)]
    return lines


def parse_test_invoice(lines):
    supplier = ""
    invoice_no = ""
    invoice_date = ""

    for line in lines:
        if line.lower().startswith("supplier:"):
            supplier = clean(line.split(":", 1)[1])
        elif line.lower().startswith("invoice no:"):
            invoice_no = clean(line.split(":", 1)[1])
        elif line.lower().startswith("invoice date:"):
            invoice_date = clean(line.split(":", 1)[1])

    try:
        start = next(i for i, x in enumerate(lines)
                     if x.lower() == "material code") + 1
    except StopIteration:
        raise ValueError("Could not find the invoice item header.")

    items = []
    i = start
    stop_words = {"subtotal", "vat 23%", "total"}

    # Test invoice extraction is 5 lines per item:
    # code, description, qty, unit price, line total
    while i + 4 < len(lines):
        if lines[i].lower() in stop_words:
            break

        code = lines[i]
        desc = lines[i + 1]

        if lines[i + 2].lower() in stop_words:
            break

        qty_text = lines[i + 2]
        unit_text = lines[i + 3]
        total_text = lines[i + 4]

        try:
            qty = float(qty_text)
            unit_price = money(unit_text)
            line_total = money(total_text)
        except Exception:
            # If this block is not a valid item, move one line forward.
            i += 1
            continue

        items.append({
            "material_id": "" if code.upper() == "N/A" else code,
            "material": desc,
            "quantity": qty,
            "unit_price": unit_price,
            "line_total": line_total,
        })
        i += 5

    return {
        "supplier": supplier,
        "invoice_no": invoice_no,
        "invoice_date": invoice_date,
        "items": items,
    }


def headers(ws):
    return {clean(c.value): c.column for c in ws[1] if c.value is not None}


def ensure_sheet(wb, name, columns):
    if name not in wb.sheetnames:
        ws = wb.create_sheet(name)
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9EAF7")
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        return ws
    return wb[name]


def setup_price_sheet(wb):
    if "MATERIAL_PRICES" not in wb.sheetnames:
        ws = wb.create_sheet("MATERIAL_PRICES")
        ws.append([
            "Material ID", "Material", "Unit", "Supplier",
            "New Price (€)", "Previous Price (€)", "Difference (€)",
            "Change %", "Last Invoice", "Last Updated"
        ])

        # Build the initial material list from the existing STOCK sheet.
        stock = wb["STOCK"]
        for row in stock.iter_rows(min_row=2, values_only=True):
            mid, material = row[0], row[1]
            if material:
                ws.append([mid or "", material, "", "", 0, 0,
                           f"=E{ws.max_row+1}-F{ws.max_row+1}",
                           f'=IF(F{ws.max_row+1}=0,"",G{ws.max_row+1}/F{ws.max_row+1})',
                           "", ""])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9EAF7")
        ws.freeze_panes = "A2"
    return wb["MATERIAL_PRICES"]


def find_material(price_ws, item):
    h = headers(price_ws)
    id_col = h.get("Material ID")
    name_col = h.get("Material")

    item_id = clean(item["material_id"])
    item_name = norm(item["material"])

    # 1. Exact Material ID
    if item_id and id_col:
        for r in range(2, price_ws.max_row + 1):
            value = clean(price_ws.cell(r, id_col).value)
            if value and value == item_id:
                return r, "ID"

    # 2. Exact normalized material name
    if name_col:
        for r in range(2, price_ws.max_row + 1):
            value = norm(price_ws.cell(r, name_col).value)
            if value and value == item_name:
                return r, "NAME"

    return None, None


def main():
    if len(sys.argv) != 3:
        print("Usage:")
        print("  python3 invoice_pdf_import.py INVOICE.pdf WORKBOOK.xlsx")
        raise SystemExit(1)

    pdf_file = Path(sys.argv[1])
    workbook_file = Path(sys.argv[2])

    if not pdf_file.exists():
        raise FileNotFoundError(pdf_file)
    if not workbook_file.exists():
        raise FileNotFoundError(workbook_file)

    lines = read_pdf(pdf_file)
    invoice = parse_test_invoice(lines)

    print(f"\nInvoice:  {invoice['invoice_no']}")
    print(f"Supplier: {invoice['supplier']}")
    print(f"Date:     {invoice['invoice_date']}")
    print(f"Items:    {len(invoice['items'])}\n")

    wb = load_workbook(workbook_file)
    price_ws = setup_price_sheet(wb)

    inv_ws = ensure_sheet(wb, "INVOICES",
                          ["Invoice No.", "Supplier", "Invoice Date",
                           "PDF File", "Invoice Total (€)", "Checked"])
    item_ws = ensure_sheet(wb, "INVOICE_ITEMS",
                           ["Invoice No.", "Material ID", "Material",
                            "Quantity", "Unit Price (€)", "Line Total (€)"])
    hist_ws = ensure_sheet(wb, "PRICE_HISTORY",
                           ["Date", "Material ID", "Material", "Supplier",
                            "Price (€)", "Invoice No."])

    updated = 0
    unknown = []

    # Add invoice header if not already present.
    inv_h = headers(inv_ws)
    invoice_exists = False
    for r in range(2, inv_ws.max_row + 1):
        if clean(inv_ws.cell(r, inv_h["Invoice No."]).value) == invoice["invoice_no"]:
            invoice_exists = True
            break

    invoice_total = sum(x["line_total"] for x in invoice["items"])

    if not invoice_exists:
        inv_ws.append([
            invoice["invoice_no"], invoice["supplier"], invoice["invoice_date"],
            pdf_file.name, invoice_total, "No"
        ])

    p_h = headers(price_ws)

    for item in invoice["items"]:
        row, match_type = find_material(price_ws, item)

        item_ws.append([
            invoice["invoice_no"], item["material_id"], item["material"],
            item["quantity"], item["unit_price"], item["line_total"]
        ])

        if row is None:
            unknown.append(item)
            continue

        old_new = price_ws.cell(row, p_h["New Price (€)"]).value
        price_ws.cell(row, p_h["Previous Price (€)"]).value = old_new
        price_ws.cell(row, p_h["New Price (€)"]).value = item["unit_price"]
        price_ws.cell(row, p_h["Supplier"]).value = invoice["supplier"]
        price_ws.cell(row, p_h["Last Invoice"]).value = invoice["invoice_no"]
        price_ws.cell(row, p_h["Last Updated"]).value = invoice["invoice_date"]

        # Keep formulas for difference and percentage.
        price_ws.cell(row, p_h["Difference (€)"]).value = (
            f"=E{row}-F{row}"
        )
        price_ws.cell(row, p_h["Change %"]).value = (
            f'=IF(F{row}=0,"",G{row}/F{row})'
        )

        mid = price_ws.cell(row, p_h["Material ID"]).value
        material = price_ws.cell(row, p_h["Material"]).value

        hist_ws.append([
            invoice["invoice_date"], mid, material,
            invoice["supplier"], item["unit_price"], invoice["invoice_no"]
        ])

        print(f"✓ {material} -> €{item['unit_price']:.2f}  (matched by {match_type})")
        updated += 1

    wb.save(workbook_file)

    print("\n--------------------------------")
    print(f"Updated prices: {updated}")
    print(f"Unknown materials: {len(unknown)}")

    if unknown:
        print("\n⚠ UNKNOWN MATERIALS - REVIEW MANUALLY:")
        for x in unknown:
            print(
                f"  ID={x['material_id'] or 'N/A'} | "
                f"{x['material']} | "
                f"€{x['unit_price']:.2f}"
            )

    print(f"\nSaved: {workbook_file}")


if __name__ == "__main__":
    main()
